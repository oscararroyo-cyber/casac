from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Título principal
titulo = doc.add_heading('PREGUNTAS DE EVALUACIÓN', 0)
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitulo = doc.add_heading('Administración Económica de Unidades (EA)', 2)
subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

modulo = doc.add_paragraph('Módulo de Logística – Administración de Unidades')
modulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('')

preguntas = [
    {
        "num": 1,
        "texto": "¿Qué Real Decreto establece la estructura orgánica del Ejército del Aire y del Espacio y crea la Dirección de Asuntos Económicos (DAE)?",
        "opciones": [
            ("a", "Real Decreto 999/1980", False),
            ("b", "Real Decreto 1108/1978", True),
            ("c", "Real Decreto 1108/1980", False),
            ("d", "Real Decreto 999/1978", False),
        ],
    },
    {
        "num": 2,
        "texto": "¿Qué norma crea las Secciones Económico-Administrativas (SEA)?",
        "opciones": [
            ("a", "Real Decreto 1108/1978", False),
            ("b", "Instrucción General 10-10", False),
            ("c", "Orden Ministerial 999/1980", True),
            ("d", "Orden Ministerial 1108/1978", False),
        ],
    },
    {
        "num": 3,
        "texto": "¿De quién depende orgánicamente la SEA?",
        "opciones": [
            ("a", "De la DAE", False),
            ("b", "Del Jefe del Mando o UCO donde esté encuadrada", True),
            ("c", "Del Ministerio de Defensa", False),
            ("d", "Del Órgano de Contratación", False),
        ],
    },
    {
        "num": 4,
        "texto": "¿Qué negociado es responsable de la ejecución presupuestaria y las altas y bajas de material?",
        "opciones": [
            ("a", "Negociado de Contratación", False),
            ("b", "Negociado de Gestión de Tesorería", False),
            ("c", "Negociado de Contabilidad", True),
            ("d", "Negociado de Patrimonio", False),
        ],
    },
    {
        "num": 5,
        "texto": "¿En qué tipo de unidades está presente el Órgano de Apoyo Económico (OAE)?",
        "opciones": [
            ("a", "En todas las unidades del EA", False),
            ("b", "En unidades que ya tienen SEA propia", False),
            ("c", "En unidades sin SEA propia", True),
            ("d", "Solo en bases aéreas principales", False),
        ],
    },
    {
        "num": 6,
        "texto": "¿Cuál es el supremo órgano fiscalizador de las cuentas del Estado?",
        "opciones": [
            ("a", "La IGAE", False),
            ("b", "El Ministerio de Hacienda", False),
            ("c", "El Tribunal de Cuentas", True),
            ("d", "La DAE", False),
        ],
    },
    {
        "num": 7,
        "texto": "¿En qué momento actúa la función interventora previa?",
        "opciones": [
            ("a", "Después de ejecutar el gasto", False),
            ("b", "Durante la auditoría anual", False),
            ("c", "Antes de aprobar actos que generan derechos u obligaciones", True),
            ("d", "Al cierre del ejercicio presupuestario", False),
        ],
    },
]

for p in preguntas:
    # Número y enunciado
    enunciado = doc.add_paragraph()
    run_num = enunciado.add_run(f"Pregunta {p['num']}. ")
    run_num.bold = True
    run_num.font.size = Pt(12)
    run_texto = enunciado.add_run(p['texto'])
    run_texto.font.size = Pt(12)

    # Opciones
    for letra, texto, correcta in p['opciones']:
        opcion = doc.add_paragraph(style='List Bullet')
        run = opcion.add_run(f"{letra}) {texto}")
        run.font.size = Pt(11)
        if correcta:
            run.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x7A, 0x1F)  # verde

    doc.add_paragraph('')

# Nota al pie
nota = doc.add_paragraph('* Las respuestas correctas están marcadas en negrita y color verde.')
nota.runs[0].font.size = Pt(9)
nota.runs[0].italic = True

doc.save('/home/user/casac/Quiz_Administracion_Economica.docx')
print("Documento generado correctamente.")
